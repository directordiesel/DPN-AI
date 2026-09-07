from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Any, Iterable

from docx import Document
from openpyxl import load_workbook
from pptx import Presentation
from pypdf import PdfReader


@dataclass(frozen=True)
class ArtifactAcceptanceResult:
    path: str
    artifact_type: str
    accepted: bool
    required_items: tuple[str, ...]
    matched_items: tuple[str, ...]
    missing_items: tuple[str, ...]
    metrics: dict[str, Any]

    def to_dict(self) -> dict[str, Any]:
        return {
            "path": self.path,
            "type": self.artifact_type,
            "accepted": self.accepted,
            "required_items": list(self.required_items),
            "matched_items": list(self.matched_items),
            "missing_items": list(self.missing_items),
            "metrics": dict(self.metrics),
        }


def _normalize(values: Iterable[str]) -> tuple[str, ...]:
    cleaned: list[str] = []
    seen: set[str] = set()
    for value in values:
        item = " ".join(str(value).strip().split())
        key = item.casefold()
        if item and key not in seen:
            cleaned.append(item)
            seen.add(key)
    return tuple(cleaned)


def _contains(haystack: Iterable[str], needle: str) -> bool:
    expected = " ".join(needle.strip().split()).casefold()
    return any(expected == " ".join(item.strip().split()).casefold() for item in haystack if item.strip())


def _docx_items(target: Path) -> tuple[list[str], dict[str, Any]]:
    document = Document(target)
    values = [str(document.core_properties.title or "").strip()]
    values.extend(paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip())
    return values, {"paragraphs": len(document.paragraphs), "tables": len(document.tables)}


def _pdf_items(target: Path) -> tuple[list[str], dict[str, Any]]:
    reader = PdfReader(str(target))
    values: list[str] = []
    try:
        values.append(str((reader.metadata or {}).get("/Title") or "").strip())
    except Exception:
        values.append("")
    pages_with_text = 0
    for page in reader.pages:
        text = (page.extract_text() or "").strip()
        if text:
            pages_with_text += 1
            values.extend(line.strip() for line in text.splitlines() if line.strip())
    return values, {"pages": len(reader.pages), "pages_with_text": pages_with_text}


def _xlsx_items(target: Path) -> tuple[list[str], dict[str, Any]]:
    workbook = load_workbook(target, data_only=False, read_only=True)
    values = list(workbook.sheetnames)
    header_cells = 0
    for worksheet in workbook.worksheets:
        for cell in next(worksheet.iter_rows(min_row=1, max_row=1), ()):
            if cell.value not in (None, ""):
                values.append(str(cell.value).strip())
                header_cells += 1
    return values, {"sheets": len(workbook.worksheets), "header_cells": header_cells}


def _pptx_items(target: Path) -> tuple[list[str], dict[str, Any]]:
    presentation = Presentation(target)
    values: list[str] = []
    slides_with_text = 0
    for slide in presentation.slides:
        slide_values = [shape.text.strip() for shape in slide.shapes if hasattr(shape, "text") and shape.text.strip()]
        if slide_values:
            slides_with_text += 1
            values.extend(slide_values)
    return values, {"slides": len(presentation.slides), "slides_with_text": slides_with_text}


def inspect_artifact_acceptance(
    path: Path,
    workspace: Path,
    *,
    required_items: Iterable[str],
) -> ArtifactAcceptanceResult:
    workspace = workspace.resolve()
    target = path.resolve()
    target.relative_to(workspace)
    relative = target.relative_to(workspace).as_posix()
    required = _normalize(required_items)
    suffix = target.suffix.lower()

    if not target.exists() or not target.is_file():
        return ArtifactAcceptanceResult(relative, suffix.lstrip(".") or "unknown", False, required, (), required, {})
    if suffix not in {".docx", ".pdf", ".xlsx", ".pptx"}:
        return ArtifactAcceptanceResult(relative, suffix.lstrip(".") or "unknown", False, required, (), required, {})
    if not required:
        return ArtifactAcceptanceResult(relative, suffix.lstrip("."), False, (), (), (), {"reason": "no required acceptance items"})

    try:
        if suffix == ".docx":
            observed, metrics = _docx_items(target)
        elif suffix == ".pdf":
            observed, metrics = _pdf_items(target)
        elif suffix == ".xlsx":
            observed, metrics = _xlsx_items(target)
        else:
            observed, metrics = _pptx_items(target)
    except Exception as exc:
        return ArtifactAcceptanceResult(
            relative,
            suffix.lstrip("."),
            False,
            required,
            (),
            required,
            {"inspection_error": type(exc).__name__},
        )

    matched = tuple(item for item in required if _contains(observed, item))
    missing = tuple(item for item in required if item not in matched)
    metrics = {**metrics, "observed_items": len(observed), "required_items": len(required), "matched_items": len(matched)}
    return ArtifactAcceptanceResult(
        path=relative,
        artifact_type=suffix.lstrip("."),
        accepted=not missing and len(matched) == len(required),
        required_items=required,
        matched_items=matched,
        missing_items=missing,
        metrics=metrics,
    )


__all__ = ["ArtifactAcceptanceResult", "inspect_artifact_acceptance"]
