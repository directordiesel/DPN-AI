from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Any

from docx import Document
from openpyxl import load_workbook
from pptx import Presentation
from pypdf import PdfReader


SUPPORTED_PROFESSIONAL_FORMATS = {".docx", ".pdf", ".xlsx", ".pptx"}


@dataclass(frozen=True)
class ArtifactQualityResult:
    path: str
    artifact_type: str
    professional_ready: bool
    score: float
    checks: tuple[str, ...]
    failures: tuple[str, ...]
    metrics: dict[str, Any]

    def to_dict(self) -> dict[str, Any]:
        return {
            "path": self.path,
            "type": self.artifact_type,
            "professional_ready": self.professional_ready,
            "score": self.score,
            "checks": list(self.checks),
            "failures": list(self.failures),
            "metrics": dict(self.metrics),
        }


def _bounded_score(passed: int, total: int) -> float:
    if total <= 0:
        return 0.0
    return round(max(0.0, min(1.0, passed / total)), 4)


def _docx_quality(target: Path) -> tuple[list[str], list[str], dict[str, Any]]:
    checks: list[str] = []
    failures: list[str] = []
    document = Document(target)
    paragraphs = [p.text.strip() for p in document.paragraphs if p.text.strip()]
    heading_count = sum(1 for p in document.paragraphs if p.style and p.style.name.startswith("Heading") and p.text.strip())
    table_count = len(document.tables)
    if paragraphs:
        checks.append("document contains readable text")
    else:
        failures.append("document has no readable text")
    if document.core_properties.title and document.core_properties.title.strip():
        checks.append("document title metadata is present")
    else:
        failures.append("document title metadata is missing")
    if heading_count >= 1:
        checks.append("document contains structured heading content")
    else:
        failures.append("document contains no structured headings")
    metrics = {"paragraphs": len(paragraphs), "headings": heading_count, "tables": table_count}
    return checks, failures, metrics


def _pdf_quality(target: Path) -> tuple[list[str], list[str], dict[str, Any]]:
    checks: list[str] = []
    failures: list[str] = []
    reader = PdfReader(str(target))
    page_count = len(reader.pages)
    text_chunks: list[str] = []
    for page in reader.pages:
        try:
            text_chunks.append((page.extract_text() or "").strip())
        except Exception:
            text_chunks.append("")
    readable_pages = sum(1 for text in text_chunks if text)
    if page_count >= 1:
        checks.append("PDF contains at least one page")
    else:
        failures.append("PDF contains no pages")
    if readable_pages >= 1:
        checks.append("PDF contains extractable text")
    else:
        failures.append("PDF contains no extractable text")
    title = ""
    try:
        title = str((reader.metadata or {}).get("/Title") or "").strip()
    except Exception:
        title = ""
    if title:
        checks.append("PDF title metadata is present")
    else:
        failures.append("PDF title metadata is missing")
    metrics = {"pages": page_count, "readable_pages": readable_pages, "title_present": bool(title)}
    return checks, failures, metrics


def _xlsx_quality(target: Path) -> tuple[list[str], list[str], dict[str, Any]]:
    checks: list[str] = []
    failures: list[str] = []
    workbook = load_workbook(target, data_only=False)
    populated_cells = 0
    formula_cells = 0
    frozen_sheets = 0
    filtered_sheets = 0
    for worksheet in workbook.worksheets:
        if worksheet.freeze_panes:
            frozen_sheets += 1
        if worksheet.auto_filter.ref:
            filtered_sheets += 1
        for row in worksheet.iter_rows():
            for cell in row:
                if cell.value not in (None, ""):
                    populated_cells += 1
                    if isinstance(cell.value, str) and cell.value.startswith("="):
                        formula_cells += 1
    if workbook.worksheets:
        checks.append("workbook contains worksheets")
    else:
        failures.append("workbook contains no worksheets")
    if populated_cells >= 1:
        checks.append("workbook contains populated cells")
    else:
        failures.append("workbook contains no populated cells")
    if frozen_sheets == len(workbook.worksheets) and workbook.worksheets:
        checks.append("worksheet headers are frozen")
    else:
        failures.append("one or more worksheets do not freeze the header row")
    if filtered_sheets == len(workbook.worksheets) and workbook.worksheets:
        checks.append("worksheet filters are configured")
    else:
        failures.append("one or more worksheets do not configure filters")
    metrics = {
        "sheets": len(workbook.worksheets),
        "populated_cells": populated_cells,
        "formula_cells": formula_cells,
        "frozen_sheets": frozen_sheets,
        "filtered_sheets": filtered_sheets,
    }
    return checks, failures, metrics


def _pptx_quality(target: Path) -> tuple[list[str], list[str], dict[str, Any]]:
    checks: list[str] = []
    failures: list[str] = []
    presentation = Presentation(target)
    slide_count = len(presentation.slides)
    text_slides = 0
    for slide in presentation.slides:
        text = " ".join(shape.text.strip() for shape in slide.shapes if hasattr(shape, "text") and shape.text.strip())
        if text:
            text_slides += 1
    if slide_count >= 2:
        checks.append("presentation contains title and content slides")
    else:
        failures.append("presentation requires a title slide and at least one content slide")
    if text_slides == slide_count and slide_count:
        checks.append("all presentation slides contain readable text")
    else:
        failures.append("one or more presentation slides contain no readable text")
    widescreen_ratio = round(float(presentation.slide_width) / float(presentation.slide_height), 3)
    if 1.7 <= widescreen_ratio <= 1.8:
        checks.append("presentation uses professional widescreen dimensions")
    else:
        failures.append("presentation is not configured for widescreen output")
    metrics = {"slides": slide_count, "text_slides": text_slides, "aspect_ratio": widescreen_ratio}
    return checks, failures, metrics


def inspect_artifact_quality(path: Path, workspace: Path) -> ArtifactQualityResult:
    workspace = workspace.resolve()
    target = path.resolve()
    target.relative_to(workspace)
    suffix = target.suffix.lower()
    relative = target.relative_to(workspace).as_posix()

    if suffix not in SUPPORTED_PROFESSIONAL_FORMATS:
        return ArtifactQualityResult(relative, suffix.lstrip(".") or "unknown", False, 0.0, (), ("unsupported professional artifact format",), {})
    if not target.exists() or not target.is_file():
        return ArtifactQualityResult(relative, suffix.lstrip("."), False, 0.0, (), ("artifact file is missing",), {})

    try:
        if suffix == ".docx":
            checks, failures, metrics = _docx_quality(target)
        elif suffix == ".pdf":
            checks, failures, metrics = _pdf_quality(target)
        elif suffix == ".xlsx":
            checks, failures, metrics = _xlsx_quality(target)
        else:
            checks, failures, metrics = _pptx_quality(target)
    except Exception as exc:
        return ArtifactQualityResult(
            relative,
            suffix.lstrip("."),
            False,
            0.0,
            (),
            (f"artifact quality inspection failed: {type(exc).__name__}",),
            {},
        )

    total = len(checks) + len(failures)
    score = _bounded_score(len(checks), total)
    return ArtifactQualityResult(
        path=relative,
        artifact_type=suffix.lstrip("."),
        professional_ready=not failures and bool(checks),
        score=score,
        checks=tuple(checks),
        failures=tuple(failures),
        metrics=metrics,
    )


__all__ = ["ArtifactQualityResult", "SUPPORTED_PROFESSIONAL_FORMATS", "inspect_artifact_quality"]
